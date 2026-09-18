from pathlib import Path

from openpyxl import Workbook


def create_workbook():
    """
    Create the reporting workbook and required worksheets.
    """

    workbook = Workbook()

    executive_sheet = workbook.active
    executive_sheet.title = "Executive Summary"

    workbook.create_sheet("Monthly Performance")
    workbook.create_sheet("Category Performance")
    workbook.create_sheet("Product Performance")
    workbook.create_sheet("Customer Performance")
    workbook.create_sheet("Seller Performance")

    return workbook


def write_dataframe_to_sheet(
    workbook,
    dataframe,
    sheet_name,
):
    """
    Write a DataFrame to a worksheet starting at A1.

    Intended for supporting/detail worksheets.
    """

    worksheet = workbook[sheet_name]

    for column_number, column_name in enumerate(
        dataframe.columns,
        start=1,
    ):
        worksheet.cell(
            row=1,
            column=column_number,
            value=column_name,
        )

    for row_number, row in enumerate(
        dataframe.itertuples(index=False),
        start=2,
    ):
        for column_number, value in enumerate(
            row,
            start=1,
        ):
            worksheet.cell(
                row=row_number,
                column=column_number,
                value=value,
            )

    return worksheet


def build_executive_summary(
    workbook,
    executive_kpis_df,
    observations,
):
    """
    Build the Executive Summary page.
    """

    worksheet = workbook["Executive Summary"]

    kpis = executive_kpis_df.iloc[0]

    worksheet["A1"] = "COMMERCIAL PERFORMANCE REPORT"
    worksheet["A2"] = "May 2011 – June 2014"

    worksheet["A4"] = "REVENUE"
    worksheet["C4"] = "GROSS PROFIT"
    worksheet["E4"] = "GROSS MARGIN"

    worksheet["A5"] = kpis["Revenue"]
    worksheet["C5"] = kpis["GrossProfit"]
    worksheet["E5"] = kpis["GrossMarginPct"]

    worksheet["A7"] = "UNITS SOLD"
    worksheet["C7"] = "CUSTOMERS"
    worksheet["E7"] = "PRODUCTS"

    worksheet["A8"] = kpis["UnitsSold"]
    worksheet["C8"] = kpis["Customers"]
    worksheet["E8"] = kpis["Products"]

    worksheet["A10"] = "REVENUE TREND"
    worksheet["A25"] = "COMMERCIAL PERFORMANCE"
    worksheet["A40"] = "WHAT NEEDS ATTENTION"

    observation_row = 41

    for observation in observations:
        worksheet.cell(
            row=observation_row,
            column=1,
            value=f"• {observation}",
        )

        observation_row += 1

    return worksheet


def build_category_performance(
    workbook,
    category_performance_df,
):
    """
    Build the Category Performance page.
    """

    worksheet = workbook["Category Performance"]

    worksheet["A1"] = "CATEGORY PERFORMANCE"

    worksheet["A2"] = (
        "Revenue, profitability and margin by product category"
    )

    worksheet["A4"] = "CATEGORY PERFORMANCE SUMMARY"

    for column_number, column_name in enumerate(
        category_performance_df.columns,
        start=1,
    ):
        worksheet.cell(
            row=5,
            column=column_number,
            value=column_name,
        )

    for row_number, row in enumerate(
        category_performance_df.itertuples(index=False),
        start=6,
    ):
        for column_number, value in enumerate(
            row,
            start=1,
        ):
            worksheet.cell(
                row=row_number,
                column=column_number,
                value=value,
            )

    return worksheet


def build_product_performance(
    workbook,
    top_products_df,
    product_margin_risk_df,
):
    """
    Build the Product Performance management page.

    The page intentionally uses management-level product
    views rather than displaying all 266 products.
    """

    worksheet = workbook["Product Performance"]

    worksheet["A1"] = "PRODUCT PERFORMANCE"

    worksheet["A2"] = (
        "Revenue leaders and commercially material margin risk"
    )

    # ---------------------------------------------------------
    # Top products section
    # ---------------------------------------------------------

    worksheet["A4"] = "TOP PRODUCTS BY REVENUE"

    top_product_columns = [
        "ProductID",
        "ProductName",
        "ProductCategory",
        "ProductSubcategory",
        "Revenue",
        "GrossProfit",
        "Units",
        "GrossMarginPct",
    ]

    for column_number, column_name in enumerate(
        top_product_columns,
        start=1,
    ):
        worksheet.cell(
            row=5,
            column=column_number,
            value=column_name,
        )

    top_products = top_products_df[
        top_product_columns
    ]

    for row_number, row in enumerate(
        top_products.itertuples(index=False),
        start=6,
    ):
        for column_number, value in enumerate(
            row,
            start=1,
        ):
            worksheet.cell(
                row=row_number,
                column=column_number,
                value=value,
            )

    # ---------------------------------------------------------
    # Margin-risk section
    # ---------------------------------------------------------

    risk_start_row = 23

    worksheet.cell(
        row=risk_start_row,
        column=1,
        value="LOW-MARGIN PRODUCTS — TOP 50 REVENUE PRODUCTS",
    )

    risk_header_row = risk_start_row + 1

    for column_number, column_name in enumerate(
        top_product_columns,
        start=1,
    ):
        worksheet.cell(
            row=risk_header_row,
            column=column_number,
            value=column_name,
        )

    for row_number, row in enumerate(
        product_margin_risk_df[
            top_product_columns
        ].itertuples(index=False),
        start=risk_header_row + 1,
    ):
        for column_number, value in enumerate(
            row,
            start=1,
        ):
            worksheet.cell(
                row=row_number,
                column=column_number,
                value=value,
            )

    return worksheet


def build_report_workbook(datasets):
    """
    Build the complete reporting workbook.
    """

    workbook = create_workbook()

    build_executive_summary(
        workbook,
        datasets["executive_kpis"],
        datasets["key_observations"],
    )

    write_dataframe_to_sheet(
        workbook,
        datasets["monthly_performance"],
        "Monthly Performance",
    )

    build_category_performance(
        workbook,
        datasets["category_performance"],
    )

    build_product_performance(
        workbook,
        datasets["top_products"],
        datasets["product_margin_risk"],
    )

    write_dataframe_to_sheet(
        workbook,
        datasets["customer_performance"],
        "Customer Performance",
    )

    write_dataframe_to_sheet(
        workbook,
        datasets["seller_performance"],
        "Seller Performance",
    )

    return workbook


def save_workbook(
    workbook,
    output_path,
):
    """
    Save the workbook to disk.
    """

    output_path = Path(output_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook.save(output_path)

    return output_path