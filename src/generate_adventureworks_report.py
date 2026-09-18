# ============================================================
# Imports
# ============================================================

from pathlib import Path

import pandas as pd

from src.adventureworks.connection import create_database_engine

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference

# ============================================================
# Configuration
# ============================================================

SERVER = r"localhost\SQLEXPRESS"
DATABASE = "AdventureWorks2022"

REPORT_DIR = Path("reports")

REPORT_PATH = (
    REPORT_DIR
    / "adventureworks_commercial_performance_report.xlsx"
)


# ============================================================
# Sales extraction
# ============================================================

def extract_sales_detail(engine):
    """
    Extract the AdventureWorks sales transaction grain.

    Grain:
        One row per SalesOrderDetailID.
    """

    query = """
        SELECT
            soh.SalesOrderID,
            sod.SalesOrderDetailID,
            soh.OrderDate,
            soh.CustomerID,
            soh.SalesPersonID,
            soh.TerritoryID,
            sod.ProductID,
            sod.OrderQty,
            sod.UnitPrice,
            sod.UnitPriceDiscount,
            sod.LineTotal
        FROM Sales.SalesOrderHeader AS soh
        INNER JOIN Sales.SalesOrderDetail AS sod
            ON soh.SalesOrderID = sod.SalesOrderID;
    """

    sales_detail_df = pd.read_sql_query(
        query,
        engine
    )

    return sales_detail_df


# ============================================================
# Commercial data enrichment
# ============================================================

def extract_commercial_data(engine):
    """
    Extract sales transactions with commercial dimensions.

    Grain:
        One row per SalesOrderDetailID.
    """

    query = """
        SELECT
            soh.SalesOrderID,
            sod.SalesOrderDetailID,
            soh.OrderDate,
            soh.CustomerID,
            c.AccountNumber,
            cp.FirstName,
            cp.LastName,
            soh.SalesPersonID,
            sp.FirstName AS SalesPersonFirstName,
            sp.LastName AS SalesPersonLastName,
            soh.TerritoryID,
            st.Name AS TerritoryName,
            st.CountryRegionCode,
            st.[Group] AS TerritoryGroup,
            sod.ProductID,
            prod.Name AS ProductName,
            ps.Name AS ProductSubcategory,
            pc.Name AS ProductCategory,
            sod.OrderQty,
            sod.UnitPrice,
            sod.UnitPriceDiscount,
            sod.LineTotal AS Revenue
        FROM Sales.SalesOrderHeader AS soh
        INNER JOIN Sales.SalesOrderDetail AS sod
            ON soh.SalesOrderID = sod.SalesOrderID
        INNER JOIN Sales.Customer AS c
            ON soh.CustomerID = c.CustomerID
        LEFT JOIN Person.Person AS cp
            ON c.PersonID = cp.BusinessEntityID
        LEFT JOIN Sales.SalesPerson AS sp_base
            ON soh.SalesPersonID = sp_base.BusinessEntityID
        LEFT JOIN Person.Person AS sp
            ON sp_base.BusinessEntityID = sp.BusinessEntityID
        INNER JOIN Sales.SalesTerritory AS st
            ON soh.TerritoryID = st.TerritoryID
        INNER JOIN Production.Product AS prod
            ON sod.ProductID = prod.ProductID
        LEFT JOIN Production.ProductSubcategory AS ps
            ON prod.ProductSubcategoryID = ps.ProductSubcategoryID
        LEFT JOIN Production.ProductCategory AS pc
            ON ps.ProductCategoryID = pc.ProductCategoryID;
    """

    commercial_data_df = pd.read_sql_query(
        query,
        engine
    )

    return commercial_data_df


# ============================================================
# Historical cost extraction
# ============================================================

def extract_historical_costs(engine):
    """
    Extract historical product costs from AdventureWorks.
    """

    query = """
        SELECT
            sod.SalesOrderDetailID,
            sod.ProductID,
            soh.OrderDate,
            pch.StandardCost AS UnitCost
        FROM Sales.SalesOrderHeader AS soh
        INNER JOIN Sales.SalesOrderDetail AS sod
            ON soh.SalesOrderID = sod.SalesOrderID
        LEFT JOIN Production.ProductCostHistory AS pch
            ON sod.ProductID = pch.ProductID
            AND soh.OrderDate >= pch.StartDate
            AND (
                pch.EndDate IS NULL
                OR soh.OrderDate <= pch.EndDate
            );
    """

    historical_cost_df = pd.read_sql_query(
        query,
        engine
    )

    return historical_cost_df


# ============================================================
# Current standard-cost fallback
# ============================================================

def extract_current_product_costs(engine):
    """
    Extract current standard cost for each product.

    Used only when no applicable historical
    cost record exists.
    """

    query = """
        SELECT
            ProductID,
            StandardCost AS CurrentStandardCost
        FROM Production.Product;
    """

    current_cost_df = pd.read_sql_query(
        query,
        engine
    )

    return current_cost_df


# ============================================================
# Commercial cost calculations
# ============================================================

def calculate_commercial_metrics(
    commercial_data_df,
    historical_cost_df,
    current_cost_df
):
    """
    Calculate Cost, Gross Profit and Gross Margin.
    """

    reporting_df = commercial_data_df.merge(
        historical_cost_df[
            [
                "SalesOrderDetailID",
                "UnitCost"
            ]
        ],
        on="SalesOrderDetailID",
        how="left"
    )

    reporting_df = reporting_df.merge(
        current_cost_df,
        on="ProductID",
        how="left"
    )

    reporting_df["ReportingUnitCost"] = (
        reporting_df["UnitCost"]
        .fillna(
            reporting_df["CurrentStandardCost"]
        )
    )

    reporting_df["Cost"] = (
        reporting_df["OrderQty"]
        * reporting_df["ReportingUnitCost"]
    )

    reporting_df["GrossProfit"] = (
        reporting_df["Revenue"]
        - reporting_df["Cost"]
    )

    reporting_df["GrossMarginPct"] = (
        reporting_df["GrossProfit"]
        / reporting_df["Revenue"]
    ) * 100

    return reporting_df


# ============================================================
# Financial calculation validation
# ============================================================

def validate_financial_calculations(reporting_df):
    """
    Validate the core financial calculations used by the report.
    """

    total_revenue = reporting_df["Revenue"].sum()
    total_cost = reporting_df["Cost"].sum()
    total_gross_profit = reporting_df["GrossProfit"].sum()

    expected_gross_profit = (
        total_revenue
        - total_cost
    )

    gross_profit_difference = abs(
        total_gross_profit
        - expected_gross_profit
    )

    unresolved_costs = (
        reporting_df["ReportingUnitCost"]
        .isna()
        .sum()
    )

    fallback_count = (
        reporting_df["UnitCost"]
        .isna()
        .sum()
    )

    print()
    print("FINANCIAL VALIDATION")
    print("-" * 40)

    print(
        "Revenue:",
        total_revenue
    )

    print(
        "Cost:",
        total_cost
    )

    print(
        "Gross Profit:",
        total_gross_profit
    )

    print(
        "Gross Profit reconciliation difference:",
        gross_profit_difference
    )

    print(
        "Fallback cost rows:",
        fallback_count
    )

    print(
        "Unresolved cost rows:",
        unresolved_costs
    )

    assert gross_profit_difference < 0.01, (
        "Gross Profit reconciliation failed."
    )

    assert unresolved_costs == 0, (
        "Unresolved product costs detected."
    )

    print("Financial validation: PASS")

# ============================================================
# Build reporting datasets
# ============================================================

def build_report_data(reporting_df):
    """
    Build all datasets required by the commercial performance
    report.
    """

    # --------------------------------------------------------
    # Executive KPIs
    # --------------------------------------------------------

    revenue = reporting_df["Revenue"].sum()
    cost = reporting_df["Cost"].sum()
    gross_profit = reporting_df["GrossProfit"].sum()

    executive_kpis = {
        "report_start_date": reporting_df["OrderDate"].min(),
        "report_end_date": reporting_df["OrderDate"].max(),
        "revenue": revenue,
        "cost": cost,
        "gross_profit": gross_profit,
        "gross_margin_pct": (
            gross_profit / revenue
        ) * 100,
        "sales_lines": reporting_df[
            "SalesOrderDetailID"
        ].nunique(),
        "units_sold": reporting_df[
            "OrderQty"
        ].sum(),
        "customers": reporting_df[
            "CustomerID"
        ].nunique(),
        "products": reporting_df[
            "ProductID"
        ].nunique(),
    }

    # --------------------------------------------------------
    # Monthly performance
    # --------------------------------------------------------

    reporting_df = reporting_df.copy()

    reporting_df["OrderMonth"] = (
        reporting_df["OrderDate"]
        .dt.to_period("M")
    )

    monthly_performance_df = (
        reporting_df
        .groupby("OrderMonth")
        .agg(
            Revenue=("Revenue", "sum"),
            Cost=("Cost", "sum"),
            GrossProfit=("GrossProfit", "sum"),
            SalesLines=("SalesOrderDetailID", "count")
        )
        .reset_index()
    )

    monthly_performance_df["GrossMarginPct"] = (
        monthly_performance_df["GrossProfit"]
        / monthly_performance_df["Revenue"]
    ) * 100

    # Excel cannot directly store pandas Period objects.
    monthly_performance_df["OrderMonth"] = (
        monthly_performance_df["OrderMonth"]
        .astype(str)
    )

    # --------------------------------------------------------
    # Category performance
    # --------------------------------------------------------

    category_performance_df = (
        reporting_df
        .groupby("ProductCategory")
        .agg(
            Revenue=("Revenue", "sum"),
            Cost=("Cost", "sum"),
            GrossProfit=("GrossProfit", "sum"),
            UnitsSold=("OrderQty", "sum"),
            SalesLines=("SalesOrderDetailID", "count"),
            Products=("ProductID", "nunique")
        )
        .reset_index()
    )

    category_performance_df["GrossMarginPct"] = (
        category_performance_df["GrossProfit"]
        / category_performance_df["Revenue"]
    ) * 100

    category_performance_df = (
        category_performance_df
        .sort_values("Revenue", ascending=False)
    )

    # --------------------------------------------------------
    # Product performance
    # --------------------------------------------------------

    product_performance_df = (
        reporting_df
        .groupby(
            ["ProductID", "ProductName"]
        )
        .agg(
            Revenue=("Revenue", "sum"),
            Cost=("Cost", "sum"),
            GrossProfit=("GrossProfit", "sum"),
            UnitsSold=("OrderQty", "sum"),
            SalesLines=("SalesOrderDetailID", "count")
        )
        .reset_index()
    )

    product_performance_df["GrossMarginPct"] = (
        product_performance_df["GrossProfit"]
        / product_performance_df["Revenue"]
    ) * 100

    product_performance_df = (
        product_performance_df
        .sort_values("Revenue", ascending=False)
    )

    # --------------------------------------------------------
    # Customer performance
    # --------------------------------------------------------

    customer_performance_df = (
        reporting_df
        .groupby(
            [
                "CustomerID",
                "AccountNumber",
                "FirstName",
                "LastName"
            ]
        )
        .agg(
            Revenue=("Revenue", "sum"),
            Cost=("Cost", "sum"),
            GrossProfit=("GrossProfit", "sum"),
            UnitsSold=("OrderQty", "sum"),
            SalesLines=("SalesOrderDetailID", "count"),
            ProductsPurchased=("ProductID", "nunique")
        )
        .reset_index()
    )

    customer_performance_df["GrossMarginPct"] = (
        customer_performance_df["GrossProfit"]
        / customer_performance_df["Revenue"]
    ) * 100

    customer_performance_df = (
        customer_performance_df
        .sort_values("Revenue", ascending=False)
    )

    # --------------------------------------------------------
    # Seller performance
    # --------------------------------------------------------

    seller_performance_df = (
        reporting_df
        .groupby(
            [
                "SalesPersonID",
                "SalesPersonFirstName",
                "SalesPersonLastName"
            ]
        )
        .agg(
            Revenue=("Revenue", "sum"),
            Cost=("Cost", "sum"),
            GrossProfit=("GrossProfit", "sum"),
            UnitsSold=("OrderQty", "sum"),
            SalesLines=("SalesOrderDetailID", "count"),
            Customers=("CustomerID", "nunique")
        )
        .reset_index()
    )

    seller_performance_df["GrossMarginPct"] = (
        seller_performance_df["GrossProfit"]
        / seller_performance_df["Revenue"]
    ) * 100

    seller_performance_df = (
        seller_performance_df
        .sort_values("Revenue", ascending=False)
    )

    # --------------------------------------------------------
    # Package all report datasets
    # --------------------------------------------------------

    report_data = {
        "executive": executive_kpis,
        "monthly": monthly_performance_df,
        "category": category_performance_df,
        "product": product_performance_df,
        "customer": customer_performance_df,
        "seller": seller_performance_df,
    }

    return report_data

# ============================================================
# Excel DataFrame Writer
# ============================================================

def write_dataframe_to_sheet(
    workbook,
    sheet_name,
    dataframe,
    title=None
):
    """
    Write a pandas DataFrame into an Excel worksheet.
    """

    if sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
    else:
        worksheet = workbook.create_sheet(
            sheet_name
        )

    # Clear existing worksheet contents
    for row in worksheet.iter_rows():
        for cell in row:
            cell.value = None

    current_row = 1

    # --------------------------------------------------------
    # Title
    # --------------------------------------------------------

    if title:
        worksheet.cell(
            row=current_row,
            column=1,
            value=title
        )

        worksheet.cell(
            row=current_row,
            column=1
        ).font = Font(
            bold=True,
            size=14
        )

        current_row += 2

    # --------------------------------------------------------
    # Column headers
    # --------------------------------------------------------

    for column_index, column_name in enumerate(
        dataframe.columns,
        start=1
    ):
        cell = worksheet.cell(
            row=current_row,
            column=column_index,
            value=column_name
        )

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    current_row += 1

    # --------------------------------------------------------
    # Data rows
    # --------------------------------------------------------

    for row in dataframe.itertuples(
        index=False,
        name=None
    ):
        for column_index, value in enumerate(
            row,
            start=1
        ):
            worksheet.cell(
                row=current_row,
                column=column_index,
                value=value
            )

        current_row += 1

    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    for column_index in range(
        1,
        len(dataframe.columns) + 1
    ):
        column_letter = get_column_letter(
            column_index
        )

        worksheet.column_dimensions[
            column_letter
        ].width = 18

    # --------------------------------------------------------
    # Freeze header
    # --------------------------------------------------------

    worksheet.freeze_panes = (
        f"A{3 if title else 2}"
    )

    return worksheet

# ============================================================
# Detail worksheet formatting
# ============================================================

thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)

def format_detail_sheet(
    worksheet,
    currency_columns=None,
    percentage_columns=None,
    integer_columns=None
):
    """
    Apply consistent presentation formatting to a report worksheet.
    """

    currency_columns = currency_columns or []
    percentage_columns = percentage_columns or []
    integer_columns = integer_columns or []

    # --------------------------------------------------------
    # Title
    # --------------------------------------------------------

    worksheet["A1"].font = Font(
        bold=True,
        size=14
    )

    # --------------------------------------------------------
    # Header row
    # --------------------------------------------------------

    header_row = 3

    for cell in worksheet[header_row]:
        if cell.value is not None:
            cell.font = Font(
                bold=True
            )
            cell.alignment = Alignment(
                horizontal="center"
            )
            cell.border = thin_border

    # --------------------------------------------------------
    # Data borders
    # --------------------------------------------------------

    for row in worksheet.iter_rows(
        min_row=header_row + 1
    ):
        for cell in row:
            if cell.value is not None:
                cell.border = thin_border

    # --------------------------------------------------------
    # Identify columns
    # --------------------------------------------------------

    header_map = {
        cell.value: cell.column
        for cell in worksheet[header_row]
        if cell.value is not None
    }

    # --------------------------------------------------------
    # Currency formatting
    # --------------------------------------------------------

    for column_name in currency_columns:

        if column_name in header_map:

            column_index = header_map[column_name]

            for row in range(
                header_row + 1,
                worksheet.max_row + 1
            ):
                worksheet.cell(
                    row=row,
                    column=column_index
                ).number_format = "$#,##0.00"

    # --------------------------------------------------------
    # Percentage formatting
    # --------------------------------------------------------

    for column_name in percentage_columns:

        if column_name in header_map:

            column_index = header_map[column_name]

            for row in range(
                header_row + 1,
                worksheet.max_row + 1
            ):
                cell = worksheet.cell(
                    row=row,
                    column=column_index
                )

                if isinstance(
                    cell.value,
                    (int, float)
                ):
                    cell.value = (
                        cell.value / 100
                    )

                cell.number_format = "0.00%"

    # --------------------------------------------------------
    # Integer formatting
    # --------------------------------------------------------

    for column_name in integer_columns:

        if column_name in header_map:

            column_index = header_map[column_name]

            for row in range(
                header_row + 1,
                worksheet.max_row + 1
            ):
                worksheet.cell(
                    row=row,
                    column=column_index
                ).number_format = "#,##0"

    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    for column_index in range(
        1,
        worksheet.max_column + 1
    ):

        column_letter = get_column_letter(
            column_index
        )

        max_length = 0

        for row in worksheet.iter_rows(
            min_col=column_index,
            max_col=column_index
        ):
            for cell in row:

                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 12),
            30
        )

    # --------------------------------------------------------
    # Freeze panes
    # --------------------------------------------------------

    worksheet.freeze_panes = "A4"

    # --------------------------------------------------------
    # Excel filter
    # --------------------------------------------------------

    worksheet.auto_filter.ref = (
        f"A3:"
        f"{get_column_letter(worksheet.max_column)}"
        f"{worksheet.max_row}"
    )

    return worksheet

def validate_generated_workbook(report_path):
    """
    Validate the saved Excel report.

    The workbook is reopened from disk so that validation
    confirms the actual deliverable, not only the in-memory object.
    """

    from openpyxl import load_workbook

    print()
    print("=" * 60)
    print("FINAL WORKBOOK VALIDATION")
    print("=" * 60)

    assert report_path.exists(), (
        "Report file does not exist."
    )

    saved_workbook = load_workbook(report_path)

    expected_sheets = [
        "Executive Summary",
        "Monthly Performance",
        "Category Performance",
        "Product Performance",
        "Customer Performance",
        "Seller Performance"
    ]

    actual_sheets = saved_workbook.sheetnames

    assert actual_sheets == expected_sheets, (
        f"Worksheet validation failed.\n"
        f"Expected: {expected_sheets}\n"
        f"Actual: {actual_sheets}"
    )

    print("Worksheet structure: PASS")

    executive_sheet = saved_workbook[
        "Executive Summary"
    ]

    monthly_sheet = saved_workbook[
        "Monthly Performance"
    ]

    category_sheet = saved_workbook[
        "Category Performance"
    ]

    product_sheet = saved_workbook[
        "Product Performance"
    ]

    customer_sheet = saved_workbook[
        "Customer Performance"
    ]

    seller_sheet = saved_workbook[
        "Seller Performance"
    ]

    # --------------------------------------------------------
    # Executive Summary
    # --------------------------------------------------------

    assert executive_sheet["A1"].value == (
        "AUTOMATED COMMERCIAL PERFORMANCE REPORT"
    )

    assert executive_sheet["A3"].value == (
        "Reporting Period"
    )

    assert abs(
        executive_sheet["B5"].value
        - 109846381.39988798
    ) < 0.01

    assert abs(
        executive_sheet["B8"].value
        - 0.1143212952402345
    ) < 0.000001

    print("Executive Summary: PASS")

    # --------------------------------------------------------
    # Monthly Performance
    # --------------------------------------------------------

    monthly_data_rows = (
        monthly_sheet.max_row - 3
    )

    assert monthly_data_rows == 38, (
        f"Expected 38 monthly rows, "
        f"found {monthly_data_rows}"
    )

    assert monthly_sheet["A4"].value == "2011-05"
    assert monthly_sheet["A41"].value == "2014-06"

    assert (
        monthly_sheet["B4"].number_format
        == "$#,##0.00"
    )

    
    monthly_header_map = {
        cell.value: cell.column
        for cell in monthly_sheet[3]
        if cell.value is not None
    }

    gross_margin_column = monthly_header_map[
        "GrossMarginPct"
    ]

    gross_margin_cell = monthly_sheet.cell(
        row=4,
        column=gross_margin_column
    )

    assert gross_margin_cell.number_format == "0.00%", (
        "GrossMarginPct formatting validation failed."
    )

    print(
        "Monthly Gross Margin format:",
        gross_margin_cell.number_format
    )

    print("Monthly Performance: PASS")

    # --------------------------------------------------------
    # Category Performance
    # --------------------------------------------------------

    category_data_rows = (
        category_sheet.max_row - 3
    )

    assert category_data_rows == 4

    print("Category Performance: PASS")

    # --------------------------------------------------------
    # Product Performance
    # --------------------------------------------------------

    product_data_rows = (
        product_sheet.max_row - 3
    )

    assert product_data_rows == 15

    print("Product Performance: PASS")

    # --------------------------------------------------------
    # Customer Performance
    # --------------------------------------------------------

    customer_data_rows = (
        customer_sheet.max_row - 3
    )

    assert customer_data_rows == 15

    print("Customer Performance: PASS")

    # --------------------------------------------------------
    # Seller Performance
    # --------------------------------------------------------

    seller_data_rows = (
        seller_sheet.max_row - 3
    )

    assert seller_data_rows == 17

    print("Seller Performance: PASS")

    # --------------------------------------------------------
    # Executive Chart
    # --------------------------------------------------------

    assert len(
        executive_sheet._charts
    ) == 1

    chart_title = (
        executive_sheet._charts[0]
        .title
        .tx
        .rich
        .p[0]
        .r[0]
        .t
    )

    assert chart_title == (
        "Monthly Revenue & Gross Profit Trend"
    )

    print("Executive trend chart: PASS")

    saved_workbook.close()

    print()
    print("FINAL WORKBOOK VALIDATION: PASS")

# ============================================================
# Executive Summary
# ============================================================

def create_executive_summary(
    workbook,
    executive_kpis
):
    """
    Create the Executive Summary worksheet.
    """

    worksheet = workbook.create_sheet(
        "Executive Summary"
    )

    worksheet["A1"] = (
        "AUTOMATED COMMERCIAL PERFORMANCE REPORT"
    )

    worksheet["A1"].font = Font(
        bold=True,
        size=18
    )

    worksheet["A3"] = "Reporting Period"

    worksheet["A3"].font = Font(
        bold=True
    )

    worksheet["B3"] = (
        f"{executive_kpis['report_start_date'].strftime('%Y-%m-%d')}"
        f" to "
        f"{executive_kpis['report_end_date'].strftime('%Y-%m-%d')}"
    )

    # --------------------------------------------------------
    # KPI block
    # --------------------------------------------------------

    worksheet["A5"] = "Revenue"
    worksheet["B5"] = executive_kpis["revenue"]

    worksheet["A6"] = "Cost"
    worksheet["B6"] = executive_kpis["cost"]

    worksheet["A7"] = "Gross Profit"
    worksheet["B7"] = executive_kpis["gross_profit"]

    worksheet["A8"] = "Gross Margin %"
    worksheet["B8"] = (
        executive_kpis["gross_margin_pct"]
        / 100
    )

    worksheet["A9"] = "Sales Lines"
    worksheet["B9"] = executive_kpis["sales_lines"]

    worksheet["A10"] = "Units Sold"
    worksheet["B10"] = executive_kpis["units_sold"]

    worksheet["A11"] = "Customers"
    worksheet["B11"] = executive_kpis["customers"]

    worksheet["A12"] = "Products"
    worksheet["B12"] = executive_kpis["products"]

    # --------------------------------------------------------
    # KPI formatting
    # --------------------------------------------------------

    for row in range(5, 13):
        worksheet.cell(
            row=row,
            column=1
        ).font = Font(
            bold=True
        )

    for cell in ["B5", "B6", "B7"]:
        worksheet[cell].number_format = (
            '$#,##0.00'
        )

    worksheet["B8"].number_format = "0.00%"

    for cell in ["B9", "B10", "B11", "B12"]:
        worksheet[cell].number_format = "#,##0"

    # --------------------------------------------------------
    # Borders
    # --------------------------------------------------------

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for row in range(5, 13):
        for column in range(1, 3):
            worksheet.cell(
                row=row,
                column=column
            ).border = thin_border

    # --------------------------------------------------------
    # Layout
    # --------------------------------------------------------

    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 24

    worksheet.freeze_panes = "A5"

    return worksheet

# ============================================================
# Main execution
# ============================================================

def main():

    print("=" * 60)
    print("AUTOMATED COMMERCIAL PERFORMANCE REPORT")
    print("=" * 60)

    REPORT_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    engine = create_database_engine()

    print("Database engine created successfully.")
    print(f"Database: {DATABASE}")
    print(f"Report: {REPORT_PATH}")

    sales_detail_df = extract_sales_detail(
        engine
    )

    print(
        "Sales detail extracted:",
        len(sales_detail_df),
        "rows"
    )

    commercial_data_df = extract_commercial_data(
        engine
    )

    print(
        "Commercial data extracted:",
        len(commercial_data_df),
        "rows"
    )

    print(
        "Commercial columns:",
        len(commercial_data_df.columns)
    )

    historical_cost_df = extract_historical_costs(
        engine
    )

    current_cost_df = extract_current_product_costs(
        engine
    )

    reporting_df = calculate_commercial_metrics(
        commercial_data_df,
        historical_cost_df,
        current_cost_df
    )

    validate_financial_calculations(
        reporting_df
    )

    print(
        "Historical cost records:",
        len(historical_cost_df)
    )

    print(
        "Reporting dataset:",
        len(reporting_df),
        "rows"
    )

    print(
        "Reporting columns:",
        len(reporting_df.columns)
    )

    report_data = build_report_data(
        reporting_df
    )

    # --------------------------------------------------------
    # Create workbook
    # --------------------------------------------------------

    workbook = Workbook()

    # Remove default worksheet
    if "Sheet" in workbook.sheetnames:
        del workbook["Sheet"]

    # --------------------------------------------------------
    # Executive Summary
    # --------------------------------------------------------

    create_executive_summary(
        workbook,
        report_data["executive"]
    )

    # --------------------------------------------------------
    # Supporting report worksheets
    # --------------------------------------------------------

    write_dataframe_to_sheet(
        workbook,
        "Monthly Performance",
        report_data["monthly"],
        "Monthly Commercial Performance"
    )

    write_dataframe_to_sheet(
        workbook,
        "Category Performance",
        report_data["category"],
        "Category Performance"
    )

    write_dataframe_to_sheet(
        workbook,
        "Product Performance",
        report_data["product"].head(15),
        "Top 15 Products by Revenue"
    )

    write_dataframe_to_sheet(
        workbook,
        "Customer Performance",
        report_data["customer"].head(15),
        "Top 15 Customers by Revenue"
    )

    write_dataframe_to_sheet(
        workbook,
        "Seller Performance",
        report_data["seller"],
        "Seller Performance"
    )

    # ============================================================
    # FORMAT DETAIL WORKSHEETS
    # ============================================================

    format_detail_sheet(
        workbook["Monthly Performance"],
        currency_columns=[
            "Revenue",
            "Cost",
            "GrossProfit"
        ],
        percentage_columns=[
            "GrossMarginPct"
        ],
        integer_columns=[
            "SalesLines"
        ]
    )

    format_detail_sheet(
        workbook["Category Performance"],
        currency_columns=[
            "Revenue",
            "Cost",
            "GrossProfit"
        ],
        percentage_columns=[
            "GrossMarginPct"
        ],
        integer_columns=[
            "UnitsSold",
            "SalesLines",
            "Products"
        ]
    )

    format_detail_sheet(
        workbook["Product Performance"],
        currency_columns=[
            "Revenue",
            "Cost",
            "GrossProfit"
        ],
        percentage_columns=[
            "GrossMarginPct"
        ],
        integer_columns=[
            "ProductID",
            "UnitsSold",
            "SalesLines"
        ]
    )

    format_detail_sheet(
        workbook["Customer Performance"],
        currency_columns=[
            "Revenue",
            "Cost",
            "GrossProfit"
        ],
        percentage_columns=[
            "GrossMarginPct"
        ],
        integer_columns=[
            "CustomerID",
            "UnitsSold",
            "SalesLines",
            "ProductsPurchased"
        ]
    )

    format_detail_sheet(
        workbook["Seller Performance"],
        currency_columns=[
            "Revenue",
            "Cost",
            "GrossProfit"
        ],
        percentage_columns=[
            "GrossMarginPct"
        ],
        integer_columns=[
            "SalesPersonID",
            "UnitsSold",
            "SalesLines",
            "Customers"
        ]
    )

    print("Detail worksheets formatted successfully.")

    # ============================================================
    # ADD EXECUTIVE SUMMARY TREND CHART
    # ============================================================

    monthly_sheet = workbook["Monthly Performance"]
    executive_sheet = workbook["Executive Summary"]

    trend_chart = LineChart()

    trend_chart.title = "Monthly Revenue & Gross Profit Trend"
    trend_chart.style = 2
    trend_chart.height = 8
    trend_chart.width = 18

    revenue_data = Reference(
        monthly_sheet,
        min_col=2,
        min_row=3,
        max_row=41
    )

    gross_profit_data = Reference(
        monthly_sheet,
        min_col=4,
        min_row=3,
        max_row=41
    )

    month_labels = Reference(
        monthly_sheet,
        min_col=1,
        min_row=4,
        max_row=41
    )

    trend_chart.add_data(
        revenue_data,
        titles_from_data=True
    )

    trend_chart.add_data(
        gross_profit_data,
        titles_from_data=True
        )

    trend_chart.set_categories(
        month_labels
    )

    executive_sheet.add_chart(
        trend_chart,
        "D2"
    )

    print("Executive trend chart added.")

    print()
    print("Excel workbook created.")
    print(
        "Worksheets:",
        workbook.sheetnames
    )

    # --------------------------------------------------------
    # Save workbook
    # --------------------------------------------------------

    workbook.save(REPORT_PATH)

    print()
    print("Excel report saved successfully.")
    print("File:", REPORT_PATH)

    validate_generated_workbook(REPORT_PATH)

    print()
    print("REPORT DATASETS CREATED")
    print("-" * 40)

    print(
        "Monthly rows:",
        len(report_data["monthly"])
    )

    print(
        "Category rows:",
        len(report_data["category"])
    )

    print(
        "Product rows:",
        len(report_data["product"])
    )

    print(
        "Customer rows:",
        len(report_data["customer"])
    )

    print(
        "Seller rows:",
        len(report_data["seller"])
    )

# ============================================================
# Script entry point
# ============================================================

if __name__ == "__main__":
    main()